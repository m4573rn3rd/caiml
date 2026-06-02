#!/usr/bin/env python3
# -*- coding: utf-8 -*-

"""
gt-assistant RAG (retrieval-augmented generation) layer — Phase 5.

Indexes the user's corpus (papers, conversations, attached files) into a local
ChromaDB vector store. At query time, retrieves the top-K most semantically
similar chunks and returns them with source labels so the calling code can
inject them into prompts.

Architecture:
- Embeddings: sentence-transformers all-MiniLM-L6-v2 (CPU-fast, ~80MB)
- Vector store: ChromaDB in ~/.config/gt_assistant/chroma/
- Chunking: paragraph-aware, ~500 char target, with sentence-aware fallback
- Source tracking: each chunk knows its origin (paper/conversation/attached_file)

Bridge-frame note: the system retrieves; the user curates. Indexing is
mechanical. No synthesis, no judgment about which chunks "really matter".
Source labels are preserved through retrieval so the user (and the prompt
template downstream) can distinguish considered work (papers) from working
drafts (conversations).
"""

import sys
import re
import hashlib
import argparse
from pathlib import Path
from datetime import datetime, timezone


# -------------------------------
# Paths and constants
# -------------------------------

CONF_HOME = Path.home() / ".config" / "gt_assistant"
CORPUS_DIR = CONF_HOME / "corpus"             # user drops papers here
CHROMA_DIR = CONF_HOME / "chroma"             # ChromaDB persistent store
DB_PATH = CONF_HOME / "data.db"               # conversations DB (existing)

COLLECTION_NAME = "gt_corpus"
EMBED_MODEL_NAME = "sentence-transformers/all-MiniLM-L6-v2"

# Chunking targets (in characters; the embed model is OK with overshoot)
CHUNK_TARGET_CHARS = 500
CHUNK_OVERLAP_CHARS = 80      # mild overlap reduces edge-of-chunk semantic loss
CHUNK_MAX_CHARS = 1200        # hard cap for runaway paragraphs

# Source kinds — kept short for use in chunk metadata
SOURCE_PAPER = "paper"
SOURCE_CONVERSATION = "conversation"
SOURCE_ATTACHED = "attached_file"


# -------------------------------
# Lazy-loaded singletons
# -------------------------------

_embed_model = None
_chroma_client = None
_chroma_collection = None


def get_embed_model():
    """Load sentence-transformers model on first use (takes ~3s).
    Subsequent calls return the cached instance."""
    global _embed_model
    if _embed_model is None:
        from sentence_transformers import SentenceTransformer
        _embed_model = SentenceTransformer(EMBED_MODEL_NAME)
    return _embed_model


def get_chroma_collection():
    """Open the persistent ChromaDB collection. Creates it if missing."""
    global _chroma_client, _chroma_collection
    if _chroma_collection is None:
        import chromadb
        from chromadb.config import Settings
        CHROMA_DIR.mkdir(parents=True, exist_ok=True)
        _chroma_client = chromadb.PersistentClient(
            path=str(CHROMA_DIR),
            settings=Settings(anonymized_telemetry=False),
        )
        _chroma_collection = _chroma_client.get_or_create_collection(
            name=COLLECTION_NAME,
            metadata={"hnsw:space": "cosine"},
        )
    return _chroma_collection


# -------------------------------
# Chunking
# -------------------------------

def chunk_text(text: str,
               target_chars: int = CHUNK_TARGET_CHARS,
               overlap_chars: int = CHUNK_OVERLAP_CHARS,
               max_chars: int = CHUNK_MAX_CHARS) -> list[str]:
    """Split text into ~target_chars chunks, respecting paragraph boundaries
    first and sentence boundaries second. Falls back to hard-split for very
    long paragraphs.

    Returns a list of chunk strings. Empty input returns an empty list.
    """
    if not text or not text.strip():
        return []

    # Normalize whitespace but preserve paragraph breaks
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    paragraphs = [p.strip() for p in re.split(r"\n\s*\n", text) if p.strip()]

    chunks: list[str] = []
    buf = ""
    for para in paragraphs:
        # If adding this paragraph keeps us under target, accumulate
        candidate = (buf + "\n\n" + para) if buf else para
        if len(candidate) <= target_chars:
            buf = candidate
            continue
        # Otherwise flush current buffer and start fresh with this paragraph
        if buf:
            chunks.append(buf)
            buf = ""
        # Paragraph itself is too long — split by sentences
        if len(para) > max_chars:
            for piece in _split_long_paragraph(para, target_chars, max_chars):
                chunks.append(piece)
        else:
            buf = para
    if buf:
        chunks.append(buf)

    # Apply overlap: each chunk after the first gets the tail of the previous
    if overlap_chars > 0 and len(chunks) > 1:
        overlapped = [chunks[0]]
        for prev, cur in zip(chunks, chunks[1:]):
            tail = prev[-overlap_chars:] if len(prev) > overlap_chars else prev
            overlapped.append(f"{tail}\n\n{cur}")
        chunks = overlapped

    return chunks


def _split_long_paragraph(para: str, target_chars: int, max_chars: int) -> list[str]:
    """Split a paragraph that's too long for one chunk on sentence boundaries.
    Falls back to hard-split if even the sentences are too long."""
    # Naive sentence split — good enough for English academic prose
    sentences = re.split(r"(?<=[.!?])\s+", para)
    out: list[str] = []
    buf = ""
    for sent in sentences:
        if not sent.strip():
            continue
        candidate = (buf + " " + sent) if buf else sent
        if len(candidate) <= target_chars:
            buf = candidate
        else:
            if buf:
                out.append(buf)
            # Single sentence too long? hard-split
            if len(sent) > max_chars:
                for i in range(0, len(sent), target_chars):
                    out.append(sent[i:i + target_chars])
                buf = ""
            else:
                buf = sent
    if buf:
        out.append(buf)
    return out


# -------------------------------
# Stable chunk IDs
# -------------------------------

def chunk_id(source_kind: str, source_id: str, chunk_index: int) -> str:
    """Deterministic id so re-indexing the same source replaces rather than
    duplicates chunks. Format: '{kind}:{source_id}:{index}'.

    Example: 'paper:GCSM_UnifiedFrameworkPaper_v2.docx:0'
    """
    # Keep raw — ChromaDB IDs are just strings.
    return f"{source_kind}:{source_id}:{chunk_index}"


def file_fingerprint(path: Path) -> str:
    """Hash a file's contents so we can detect changes since last index."""
    h = hashlib.sha256()
    with open(path, "rb") as f:
        for block in iter(lambda: f.read(65536), b""):
            h.update(block)
    return h.hexdigest()[:16]


# -------------------------------
# Document loaders (paper files from CORPUS_DIR)
# -------------------------------

def load_paper_text(path: Path) -> str | None:
    """Extract text from a paper file. Supports .pdf, .docx, .txt, .md.
    Returns None if extraction fails or yields no text."""
    ext = path.suffix.lower()
    try:
        if ext == ".pdf":
            from pypdf import PdfReader
            reader = PdfReader(str(path))
            if reader.is_encrypted:
                return None
            parts = []
            for i, page in enumerate(reader.pages, start=1):
                txt = page.extract_text() or ""
                if txt.strip():
                    parts.append(txt)
            return "\n\n".join(parts).strip() or None
        elif ext == ".docx":
            from docx import Document
            doc = Document(str(path))
            parts = []
            for p in doc.paragraphs:
                t = p.text.strip()
                if t:
                    parts.append(t)
            for table in doc.tables:
                for row in table.rows:
                    cells = [c.text.strip() for c in row.cells if c.text.strip()]
                    if cells:
                        parts.append(" | ".join(cells))
            return "\n\n".join(parts).strip() or None
        elif ext in (".txt", ".md"):
            return path.read_text(encoding="utf-8", errors="replace").strip() or None
        else:
            return None
    except Exception as e:
        print(f"  [warn] could not read {path.name}: {type(e).__name__}: {e}",
              file=sys.stderr)
        return None


# -------------------------------
# Indexing
# -------------------------------

def index_paper(path: Path, fingerprint: str = None) -> int:
    """Index a single paper file. Returns the number of chunks written.
    Existing chunks for this source are replaced (idempotent re-index)."""
    text = load_paper_text(path)
    if not text:
        return 0
    if fingerprint is None:
        fingerprint = file_fingerprint(path)

    source_id = path.name  # use filename as the source identifier
    chunks = chunk_text(text)
    if not chunks:
        return 0

    collection = get_chroma_collection()
    # Wipe any existing chunks for this source before re-adding
    _delete_chunks_for_source(SOURCE_PAPER, source_id)

    ids = [chunk_id(SOURCE_PAPER, source_id, i) for i in range(len(chunks))]
    embeddings = get_embed_model().encode(chunks, convert_to_numpy=True).tolist()
    metadatas = [
        {
            "source_kind": SOURCE_PAPER,
            "source_id": source_id,
            "source_path": str(path),
            "chunk_index": i,
            "fingerprint": fingerprint,
            "indexed_at": _utc_now_iso(),
        }
        for i in range(len(chunks))
    ]
    collection.add(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    return len(chunks)


def index_conversation(conn, conversation_id: int, title: str = None) -> int:
    """Index all messages in a single conversation. Concatenates user+assistant
    messages into one text stream (paragraph-separated) and chunks it.

    Idempotent: existing chunks for this conversation are replaced.
    """
    import gt_db
    messages = gt_db.get_messages(conn, conversation_id)
    if not messages:
        return 0

    # Compose a readable transcript: 'User: ...' / 'Assistant: ...' lines,
    # paragraph-separated so chunker respects message boundaries.
    parts = []
    for m in messages:
        role_label = "User" if m["role"] == "user" else "Assistant"
        ts = m.get("created_at", "")
        parts.append(f"[{role_label} at {ts}]\n{m['content']}")
    text = "\n\n".join(parts).strip()
    if not text:
        return 0

    source_id = str(conversation_id)
    chunks = chunk_text(text)
    if not chunks:
        return 0

    collection = get_chroma_collection()
    _delete_chunks_for_source(SOURCE_CONVERSATION, source_id)

    ids = [chunk_id(SOURCE_CONVERSATION, source_id, i) for i in range(len(chunks))]
    embeddings = get_embed_model().encode(chunks, convert_to_numpy=True).tolist()
    metadatas = [
        {
            "source_kind": SOURCE_CONVERSATION,
            "source_id": source_id,
            "source_title": title or "",
            "chunk_index": i,
            "indexed_at": _utc_now_iso(),
        }
        for i in range(len(chunks))
    ]
    collection.add(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    return len(chunks)


def index_attached_file(conn, file_id: int) -> int:
    """Index a single attached_files row by id. Images are skipped (b64
    content is not text-searchable). Returns chunk count written."""
    import gt_db
    record = gt_db.get_attached_file(conn, file_id)
    if not record:
        return 0
    if record["kind"] == "image":
        return 0  # skip images — no text content
    text = record["content"] or ""
    if not text.strip():
        return 0

    source_id = f"{file_id}:{record['filename']}"
    chunks = chunk_text(text)
    if not chunks:
        return 0

    collection = get_chroma_collection()
    _delete_chunks_for_source(SOURCE_ATTACHED, source_id)

    ids = [chunk_id(SOURCE_ATTACHED, source_id, i) for i in range(len(chunks))]
    embeddings = get_embed_model().encode(chunks, convert_to_numpy=True).tolist()
    metadatas = [
        {
            "source_kind": SOURCE_ATTACHED,
            "source_id": source_id,
            "source_filename": record["filename"],
            "source_conversation_id": str(record["conversation_id"]),
            "chunk_index": i,
            "indexed_at": _utc_now_iso(),
        }
        for i in range(len(chunks))
    ]
    collection.add(ids=ids, documents=chunks, embeddings=embeddings, metadatas=metadatas)
    return len(chunks)


def _delete_chunks_for_source(source_kind: str, source_id: str) -> None:
    """Remove every chunk in the collection that matches this source.
    Used for idempotent re-indexing."""
    collection = get_chroma_collection()
    try:
        collection.delete(where={"$and": [
            {"source_kind": source_kind},
            {"source_id": source_id},
        ]})
    except Exception:
        # If the where-syntax isn't supported by this chroma version, fall
        # back to listing IDs and deleting by id
        try:
            res = collection.get(where={"source_kind": source_kind})
            ids_to_delete = [
                rid for rid, md in zip(res.get("ids", []), res.get("metadatas", []))
                if md and md.get("source_id") == source_id
            ]
            if ids_to_delete:
                collection.delete(ids=ids_to_delete)
        except Exception:
            pass


# -------------------------------
# Bulk indexing entry points
# -------------------------------

def index_all_papers(verbose: bool = True) -> dict:
    """Index every supported file in CORPUS_DIR. Returns a summary dict:
        {'indexed': N, 'skipped': M, 'failed': K, 'chunks': total}
    """
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    summary = {"indexed": 0, "skipped": 0, "failed": 0, "chunks": 0}
    supported_exts = {".pdf", ".docx", ".txt", ".md"}

    paper_files = [
        p for p in sorted(CORPUS_DIR.iterdir())
        if p.is_file() and p.suffix.lower() in supported_exts
    ]
    if verbose:
        print(f"Found {len(paper_files)} paper(s) in {CORPUS_DIR}")

    for path in paper_files:
        if verbose:
            print(f"  Indexing {path.name}...", flush=True)
        try:
            fp = file_fingerprint(path)
            n = index_paper(path, fingerprint=fp)
            if n == 0:
                summary["skipped"] += 1
                if verbose:
                    print(f"    skipped (no extractable text)")
            else:
                summary["indexed"] += 1
                summary["chunks"] += n
                if verbose:
                    print(f"    {n} chunk(s)")
        except Exception as e:
            summary["failed"] += 1
            print(f"    [error] {type(e).__name__}: {e}", file=sys.stderr)

    return summary


def index_all_conversations(verbose: bool = True) -> dict:
    """Index every conversation with at least one message. Idempotent."""
    import gt_db
    conn = gt_db.ensure_schema(DB_PATH)
    gt_db.ensure_default_user(conn)
    convs = gt_db.list_conversations(conn, include_archived=True)
    summary = {"indexed": 0, "skipped": 0, "failed": 0, "chunks": 0}
    if verbose:
        print(f"Found {len(convs)} conversation(s)")
    for c in convs:
        try:
            n = index_conversation(conn, c["id"], title=c.get("title"))
            if n == 0:
                summary["skipped"] += 1
            else:
                summary["indexed"] += 1
                summary["chunks"] += n
                if verbose:
                    title = c.get("title") or "Untitled"
                    print(f"  Conv {c['id']} ({title[:40]}): {n} chunk(s)")
        except Exception as e:
            summary["failed"] += 1
            print(f"  [error] conv {c['id']}: {type(e).__name__}: {e}",
                  file=sys.stderr)
    return summary


def index_all_attached_files(verbose: bool = True) -> dict:
    """Index every text-based attached_files row (skips images)."""
    import gt_db
    conn = gt_db.ensure_schema(DB_PATH)
    gt_db.ensure_default_user(conn)
    files = gt_db.list_all_attached_files(conn, limit=10000)
    summary = {"indexed": 0, "skipped": 0, "failed": 0, "chunks": 0}
    if verbose:
        print(f"Found {len(files)} attached file(s)")
    for f in files:
        if f["kind"] == "image":
            summary["skipped"] += 1
            continue
        try:
            n = index_attached_file(conn, f["id"])
            if n == 0:
                summary["skipped"] += 1
            else:
                summary["indexed"] += 1
                summary["chunks"] += n
                if verbose:
                    print(f"  File {f['id']} ({f['filename']}): {n} chunk(s)")
        except Exception as e:
            summary["failed"] += 1
            print(f"  [error] file {f['id']}: {type(e).__name__}: {e}",
                  file=sys.stderr)
    return summary


def index_everything(verbose: bool = True) -> dict:
    """Run all three indexers. Returns a combined summary."""
    if verbose:
        print("=== Indexing papers ===")
    p = index_all_papers(verbose=verbose)
    if verbose:
        print(f"\n=== Indexing conversations ===")
    c = index_all_conversations(verbose=verbose)
    if verbose:
        print(f"\n=== Indexing attached files ===")
    a = index_all_attached_files(verbose=verbose)
    total = {
        "papers":         p,
        "conversations":  c,
        "attached_files": a,
        "total_chunks":   p["chunks"] + c["chunks"] + a["chunks"],
    }
    if verbose:
        print(f"\n=== Total: {total['total_chunks']} chunks indexed ===")
    return total


# -------------------------------
# Retrieval
# -------------------------------

def retrieve(query: str, top_k: int = 5) -> list[dict]:
    """Return the top-K most similar chunks across the entire corpus.

    Each result dict has:
        {
          'text': str,             # the chunk content
          'source_kind': str,      # 'paper' | 'conversation' | 'attached_file'
          'source_label': str,     # human-readable origin (filename / conv title / etc.)
          'distance': float,       # cosine distance, lower = more similar
          'metadata': dict,        # full chroma metadata for debugging
        }

    Empty query or empty collection returns [].
    """
    if not query or not query.strip():
        return []
    collection = get_chroma_collection()
    try:
        # If collection is empty, query will raise — guard explicitly
        if collection.count() == 0:
            return []
        embed = get_embed_model().encode([query], convert_to_numpy=True).tolist()
        res = collection.query(
            query_embeddings=embed,
            n_results=top_k,
        )
    except Exception as e:
        print(f"[retrieve] {type(e).__name__}: {e}", file=sys.stderr)
        return []

    out = []
    docs = (res.get("documents") or [[]])[0]
    metas = (res.get("metadatas") or [[]])[0]
    dists = (res.get("distances") or [[]])[0]
    for doc, meta, dist in zip(docs, metas, dists):
        out.append({
            "text": doc,
            "source_kind": meta.get("source_kind", "?"),
            "source_label": _format_source_label(meta),
            "distance": dist,
            "metadata": meta,
        })
    return out


def _format_source_label(meta: dict) -> str:
    """Human-readable origin string for a chunk.
    Different shape per source_kind to keep the most informative bit visible."""
    kind = meta.get("source_kind", "?")
    if kind == SOURCE_PAPER:
        return f"paper: {meta.get('source_id', '?')}"
    elif kind == SOURCE_CONVERSATION:
        title = meta.get("source_title") or "Untitled"
        cid = meta.get("source_id", "?")
        return f"conversation: {title[:60]} (#{cid})"
    elif kind == SOURCE_ATTACHED:
        fname = meta.get("source_filename", "?")
        cid = meta.get("source_conversation_id", "?")
        return f"attached file: {fname} (from conv #{cid})"
    return kind


def format_retrieval_block(results: list[dict]) -> str:
    """Format retrieved chunks into a 'RELEVANT CONTEXT FROM YOUR CORPUS'
    block suitable for prepending to a prompt. Source-labeled so the model
    can distinguish papers from conversations from attached files."""
    if not results:
        return ""
    lines = ["RELEVANT CONTEXT FROM YOUR CORPUS:",
             "(Mechanical retrieval — top-K by semantic similarity. Source "
             "labels preserved so you can weigh papers vs working notes.)",
             ""]
    for i, r in enumerate(results, start=1):
        lines.append(f"--- Chunk {i} [{r['source_label']}] ---")
        lines.append(r["text"])
        lines.append("")
    lines.append("END RELEVANT CONTEXT")
    return "\n".join(lines)


# -------------------------------
# Misc helpers
# -------------------------------

def _utc_now_iso() -> str:
    return datetime.now(timezone.utc).isoformat(timespec="seconds")


def collection_stats() -> dict:
    """Return summary stats about what's currently indexed.
    Useful for the eventual settings UI."""
    collection = get_chroma_collection()
    total = collection.count()
    if total == 0:
        return {"total_chunks": 0, "by_kind": {}}
    # Count by source_kind — pull metadatas in batches
    by_kind: dict[str, int] = {}
    batch = 1000
    offset = 0
    while offset < total:
        res = collection.get(limit=batch, offset=offset, include=["metadatas"])
        for md in res.get("metadatas") or []:
            kind = (md or {}).get("source_kind", "?")
            by_kind[kind] = by_kind.get(kind, 0) + 1
        offset += batch
    return {"total_chunks": total, "by_kind": by_kind}


def list_indexed_papers() -> list[dict]:
    """Return one row per indexed paper source: filename, chunk count,
    indexed_at, fingerprint. Used for the Corpus management UI.
    """
    collection = get_chroma_collection()
    total = collection.count()
    if total == 0:
        return []
    by_source: dict[str, dict] = {}
    batch = 1000
    offset = 0
    while offset < total:
        res = collection.get(
            limit=batch, offset=offset,
            include=["metadatas"],
            where={"source_kind": SOURCE_PAPER},
        )
        for md in res.get("metadatas") or []:
            if not md:
                continue
            sid = md.get("source_id", "?")
            if sid not in by_source:
                by_source[sid] = {
                    "filename": sid,
                    "source_path": md.get("source_path", ""),
                    "fingerprint": md.get("fingerprint", ""),
                    "indexed_at": md.get("indexed_at", ""),
                    "chunk_count": 0,
                }
            by_source[sid]["chunk_count"] += 1
        offset += batch
    # Stable display order: most-recently indexed first
    rows = list(by_source.values())
    rows.sort(key=lambda r: r.get("indexed_at", ""), reverse=True)
    return rows


def delete_paper(filename: str) -> int:
    """Remove all chunks for a single paper from the index. Does NOT delete
    the file from disk — that's the user's call. Returns chunks deleted."""
    collection = get_chroma_collection()
    # Count before, for the return value
    try:
        res = collection.get(where={"$and": [
            {"source_kind": SOURCE_PAPER},
            {"source_id": filename},
        ]}, include=[])
        n = len(res.get("ids", []))
    except Exception:
        # Fallback path for older chroma versions
        res = collection.get(where={"source_kind": SOURCE_PAPER}, include=["metadatas"])
        n = sum(
            1 for md in (res.get("metadatas") or [])
            if md and md.get("source_id") == filename
        )
    _delete_chunks_for_source(SOURCE_PAPER, filename)
    return n


def add_paper_from_disk(src_path: Path) -> tuple[Path, int]:
    """Copy a paper from anywhere on disk into CORPUS_DIR and index it.
    Returns (destination_path, chunks_added). If a file with the same name
    already exists in the corpus folder, it's overwritten (the indexer is
    idempotent so re-indexing replaces chunks)."""
    import shutil
    CORPUS_DIR.mkdir(parents=True, exist_ok=True)
    dst = CORPUS_DIR / src_path.name
    shutil.copy(str(src_path), str(dst))
    n = index_paper(dst)
    return dst, n


# -------------------------------
# Command-line interface (for testing in session 1)
# -------------------------------

def _cli():
    parser = argparse.ArgumentParser(
        description="gt-assistant RAG indexer / retriever",
    )
    sub = parser.add_subparsers(dest="cmd", required=True)

    sub.add_parser("index", help="Index everything (papers + conversations + attached files)")
    sub.add_parser("index-papers", help="Index only papers in ~/.config/gt_assistant/corpus")
    sub.add_parser("index-conversations", help="Index only conversations")
    sub.add_parser("index-attached", help="Index only attached files")

    q = sub.add_parser("query", help="Retrieve top-K chunks for a query")
    q.add_argument("text", nargs="+", help="The query text")
    q.add_argument("-k", "--top-k", type=int, default=5)

    sub.add_parser("stats", help="Show collection statistics")
    sub.add_parser("reset", help="Wipe the entire collection (destructive)")

    args = parser.parse_args()

    if args.cmd == "index":
        index_everything()
    elif args.cmd == "index-papers":
        s = index_all_papers()
        print(f"Done: {s}")
    elif args.cmd == "index-conversations":
        s = index_all_conversations()
        print(f"Done: {s}")
    elif args.cmd == "index-attached":
        s = index_all_attached_files()
        print(f"Done: {s}")
    elif args.cmd == "query":
        q_text = " ".join(args.text)
        print(f"Query: {q_text}\nTop-{args.top_k}:\n")
        results = retrieve(q_text, top_k=args.top_k)
        if not results:
            print("(no results — is anything indexed?)")
        for i, r in enumerate(results, start=1):
            print(f"--- {i}. [{r['source_label']}] (dist={r['distance']:.3f}) ---")
            print(r["text"][:400] + ("..." if len(r["text"]) > 400 else ""))
            print()
    elif args.cmd == "stats":
        stats = collection_stats()
        print(f"Total chunks: {stats['total_chunks']}")
        for kind, count in sorted(stats["by_kind"].items()):
            print(f"  {kind}: {count}")
    elif args.cmd == "reset":
        confirm = input("Wipe the entire collection? Type YES to confirm: ")
        if confirm.strip() == "YES":
            import chromadb
            client = chromadb.PersistentClient(path=str(CHROMA_DIR))
            try:
                client.delete_collection(COLLECTION_NAME)
                print("Collection deleted.")
            except Exception as e:
                print(f"Delete failed: {e}")
        else:
            print("Aborted.")


if __name__ == "__main__":
    _cli()
